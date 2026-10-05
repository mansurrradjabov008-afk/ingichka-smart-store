"""
data/build_real_inventory.py
Generates the real MarkazSavdo store catalog (41 verified items from video),
updates products.json, ingichka_store.db SQLite database, and creates
the professional Excel report (reports/KIYIM_KECHAK_DOKONI_OMBOR_HISOBOTI.xlsx).
"""

import os
import sys
import json
import sqlite3
from pathlib import Path

# Add project root to sys.path
ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.append(str(ROOT_DIR))
from utils.file_utils import atomic_write_json

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

REAL_PRODUCTS = [
    {
        "id": 1,
        "sku": "MS-1001",
        "name": "Poplin ayollar ko'ylak komplekt",
        "category": "Ko'ylak",
        "brand": "MarkazSavdo",
        "sizes": ["48"],
        "colors": ["to'q sariq gulli", "sariq"],
        "gender": "Ayol",
        "material": "Poplin 100%",
        "stock": 12,
        "min_stock": 2,
        "cost_price": 45000,
        "sale_price": 70000,
        "supplier": "Toshkent Tikuvchilik",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["poplin koylak", "ayollar koylagi", "gulli koylak", "poplin komplekt", "ayollar ko'ylagi", "poplin"],
        "image_url": "https://images.unsplash.com/photo-1595777457583-95e059d581b8?w=600"
    },
    {
        "id": 2,
        "sku": "MS-1002",
        "name": "Ayollar naqshli ko'ylak komplekt",
        "category": "Ko'ylak",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["jigarrang", "bej"],
        "gender": "Ayol",
        "material": "Bambuk / Trikotaj",
        "stock": 8,
        "min_stock": 2,
        "cost_price": 120000,
        "sale_price": 180000,
        "supplier": "Samarqand Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["naqshli koylak", "ayollar komplekt", "dvoyka koylak", "jigarrang koylak", "naqshli komplekt"],
        "image_url": "https://images.unsplash.com/photo-1585487000160-6ebcfceb0d03?w=600"
    },
    {
        "id": 3,
        "sku": "MS-1003",
        "name": "Qizlar keng bichimli jinsi shim",
        "category": "Shim",
        "brand": "MarkazSavdo",
        "sizes": ["Standart", "28", "30"],
        "colors": ["moviy", "och ko'k"],
        "gender": "Ayol",
        "material": "Jinsi (Denim)",
        "stock": 10,
        "min_stock": 2,
        "cost_price": 65000,
        "sale_price": 100000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["keng jinsi", "qizlar jinsisi", "moviy jinsi", "keng bichimli shim", "ayollar jinsi"],
        "image_url": "https://images.unsplash.com/photo-1541099649105-f69ad21f3246?w=600"
    },
    {
        "id": 4,
        "sku": "MS-1004",
        "name": "Polo erkaklar svitir kofta",
        "category": "Svitir",
        "brand": "Polo",
        "sizes": ["L", "XL"],
        "colors": ["kulrang", "oq"],
        "gender": "Erkak",
        "material": "Trikotaj paxta",
        "stock": 14,
        "min_stock": 3,
        "cost_price": 48000,
        "sale_price": 75000,
        "supplier": "Guangzhou Trade",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["polo svitir", "polo kofta", "erkaklar sviteri", "kulrang polo", "polo"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 5,
        "sku": "MS-1005",
        "name": "Erkaklar sportivka kostyumi",
        "category": "Kostyum",
        "brand": "MarkazSavdo",
        "sizes": ["56", "58", "60"],
        "colors": ["to'q ko'k"],
        "gender": "Erkak",
        "material": "Elastik sport trikotaj",
        "stock": 6,
        "min_stock": 2,
        "cost_price": 100000,
        "sale_price": 150000,
        "supplier": "Toshkent Sport",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["sportivka kostyum", "katta razmer sportivka", "erkaklar sportivkasi", "sport kiyim"],
        "image_url": "https://images.unsplash.com/photo-1515886657613-9f3515b0c78f?w=600"
    },
    {
        "id": 6,
        "sku": "MS-1006",
        "name": "Erkaklar yumshoq ko'k svitir",
        "category": "Svitir",
        "brand": "MarkazSavdo",
        "sizes": ["L", "XL", "2XL"],
        "colors": ["ko'k"],
        "gender": "Erkak",
        "material": "Paxta trikotaj",
        "stock": 15,
        "min_stock": 3,
        "cost_price": 32000,
        "sale_price": 50000,
        "supplier": "Buxoro Trikotaj",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["kok svitir", "yumshoq svitir", "erkaklar ko'k sviteri", "arzon svitir"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 7,
        "sku": "MS-1007",
        "name": "Erkaklar klassik qora svitir",
        "category": "Svitir",
        "brand": "MarkazSavdo",
        "sizes": ["XL", "XXL", "XXXL"],
        "colors": ["qora"],
        "gender": "Erkak",
        "material": "Trikotaj",
        "stock": 20,
        "min_stock": 3,
        "cost_price": 28000,
        "sale_price": 45000,
        "supplier": "Buxoro Trikotaj",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["qora svitir", "klassik svitir", "erkaklar qora sviteri", "arzon qora svitir"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 8,
        "sku": "MS-1008",
        "name": "Erkaklar sport triko shimi",
        "category": "Shim",
        "brand": "MarkazSavdo",
        "sizes": ["48", "50", "52", "54", "56", "58"],
        "colors": ["qora"],
        "gender": "Erkak",
        "material": "Sportiv triko",
        "stock": 18,
        "min_stock": 3,
        "cost_price": 48000,
        "sale_price": 75000,
        "supplier": "Toshkent Sport",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["sport triko", "erkaklar trikosi", "qora triko", "sport shimi"],
        "image_url": "https://images.unsplash.com/photo-1541099649105-f69ad21f3246?w=600"
    },
    {
        "id": 9,
        "sku": "MS-1009",
        "name": "Bolalar qora triko shimi",
        "category": "Shim",
        "brand": "MarkazSavdo",
        "sizes": ["48", "50"],
        "colors": ["qora"],
        "gender": "Bolalar",
        "material": "Sport triko",
        "stock": 2,
        "min_stock": 2,
        "cost_price": 15000,
        "sale_price": 25000,
        "supplier": "Toshkent Bolalar",
        "incoming_date": "04.10.2026",
        "status": "Kam qolgan",
        "aliases": ["bolalar trikosi", "qora bolalar shimi", "kichik triko"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 10,
        "sku": "MS-1010",
        "name": "Erkaklar yoqali Polo svitir",
        "category": "Svitir",
        "brand": "Polo",
        "sizes": ["L", "XL", "2XL", "3XL"],
        "colors": ["moviy"],
        "gender": "Erkak",
        "material": "Qalin paxta",
        "stock": 0,
        "min_stock": 2,
        "cost_price": 58000,
        "sale_price": 90000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Tugagan",
        "aliases": ["yoqali polo", "moviy polo svitir", "polo yoqali svitir"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 11,
        "sku": "MS-1011",
        "name": "Erkaklar tugmali to'q ko'k svitir",
        "category": "Svitir",
        "brand": "MarkazSavdo",
        "sizes": ["L", "XL", "2XL", "3XL"],
        "colors": ["to'q ko'k"],
        "gender": "Erkak",
        "material": "Paxta / Jun aralash",
        "stock": 11,
        "min_stock": 2,
        "cost_price": 55000,
        "sale_price": 85000,
        "supplier": "Samarqand Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["tugmali svitir", "to'q ko'k svitir", "tugmali kofta"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 12,
        "sku": "MS-1012",
        "name": "Ayollar charmli shippak tapichka",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["36", "37", "38", "39", "40"],
        "colors": ["qora", "jigarrang"],
        "gender": "Ayol",
        "material": "Sun'iy charm",
        "stock": 15,
        "min_stock": 3,
        "cost_price": 26000,
        "sale_price": 43000,
        "supplier": "Andijon Charm",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["charmli tapichka", "charmli shippak", "ayollar tapichkasi", "ayollar shippagi"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    },
    {
        "id": 13,
        "sku": "MS-1013",
        "name": "Ayollar guldor matoli tapichka",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["36", "37", "38", "39", "40"],
        "colors": ["qora", "bordo"],
        "gender": "Ayol",
        "material": "Guldor mato",
        "stock": 16,
        "min_stock": 3,
        "cost_price": 20000,
        "sale_price": 33000,
        "supplier": "Andijon Charm",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["guldor tapichka", "matoli tapichka", "bordo tapichka"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    },
    {
        "id": 14,
        "sku": "MS-1014",
        "name": "Ayollar naqshli moviy tapichka",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["36", "37", "38", "39", "40"],
        "colors": ["to'q ko'k", "moviy"],
        "gender": "Ayol",
        "material": "Mato",
        "stock": 12,
        "min_stock": 3,
        "cost_price": 12000,
        "sale_price": 20000,
        "supplier": "Andijon Charm",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["moviy tapichka", "naqshli tapichka", "arzon tapichka"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    },
    {
        "id": 15,
        "sku": "MS-1015",
        "name": "Erkaklar kulrang vitrofka kurtka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["50", "52", "54", "56"],
        "colors": ["kulrang"],
        "gender": "Erkak",
        "material": "Plashovka / Suv o'tkazmas",
        "stock": 10,
        "min_stock": 2,
        "cost_price": 65000,
        "sale_price": 100000,
        "supplier": "Xitoy Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["kulrang vitrofka", "vitrofka kurtka", "erkaklar vitrofkasi", "yengil kurtka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 16,
        "sku": "MS-1016",
        "name": "Erkaklar kapyushonli moviy vitrofka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["50", "52", "54", "56"],
        "colors": ["moviy", "to'q ko'k"],
        "gender": "Erkak",
        "material": "Plashovka",
        "stock": 9,
        "min_stock": 2,
        "cost_price": 90000,
        "sale_price": 140000,
        "supplier": "Xitoy Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["kapyushonli vitrofka", "moviy vitrofka", "kapyushonli kurtka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 17,
        "sku": "MS-1017",
        "name": "Erkaklar zamonaviy qora vitrofka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["50", "52", "54", "56"],
        "colors": ["qora"],
        "gender": "Erkak",
        "material": "Plashovka",
        "stock": 12,
        "min_stock": 2,
        "cost_price": 72000,
        "sale_price": 110000,
        "supplier": "Xitoy Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["qora vitrofka", "zamonaviy vitrofka", "qora kurtka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 18,
        "sku": "MS-1018",
        "name": "Erkaklar klassik qora vitrofka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["50", "52", "54", "56"],
        "colors": ["qora"],
        "gender": "Erkak",
        "material": "Sifatli plashovka",
        "stock": 8,
        "min_stock": 2,
        "cost_price": 98000,
        "sale_price": 150000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["klassik vitrofka", "qora klassik vitrofka", "sifatli kurtka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 19,
        "sku": "MS-1019",
        "name": "Chaqaloqlar konverti va Ryukzak",
        "category": "Bolalar kiyimi",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["moviy", "qora"],
        "gender": "Bolalar",
        "material": "Paxta / Sintifon",
        "stock": 5,
        "min_stock": 2,
        "cost_price": 60000,
        "sale_price": 90000,
        "supplier": "Toshkent Baby",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["chaqaloqlar konverti", "ryukzak konvert", "bolalar sumkasi", "chaqaloq konverti"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 20,
        "sku": "MS-1020",
        "name": "O'zbekiston bolalar futbolkasi",
        "category": "Futbolka",
        "brand": "MarkazSavdo",
        "sizes": ["3 yosh", "4 yosh", "5 yosh"],
        "colors": ["oq", "havorang"],
        "gender": "Bolalar",
        "material": "Paxta 100%",
        "stock": 30,
        "min_stock": 5,
        "cost_price": 5000,
        "sale_price": 8000,
        "supplier": "Namangan Paxta",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["ozbekiston futbolka", "bolalar futbolkasi", "arzon futbolka", "oq futbolka"],
        "image_url": "https://images.unsplash.com/photo-1581655353564-df123a1eb820?w=600"
    },
    {
        "id": 21,
        "sku": "MS-1021",
        "name": "Yozgi qulay tapichka shippak",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["36", "37", "38", "39", "40"],
        "colors": ["qora", "ko'k"],
        "gender": "Ayol",
        "material": "Rezina / Silikon",
        "stock": 14,
        "min_stock": 3,
        "cost_price": 20000,
        "sale_price": 33000,
        "supplier": "Andijon Charm",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["yozgi tapichka", "yozgi shippak", "silikon shippak", "rezina shippak"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    },
    {
        "id": 22,
        "sku": "MS-1022",
        "name": "5-6 yoshli o'g'il bolalar trikosi",
        "category": "Shim",
        "brand": "Adidas",
        "sizes": ["5 yosh", "6 yosh"],
        "colors": ["qora"],
        "gender": "Bolalar",
        "material": "Sport triko",
        "stock": 12,
        "min_stock": 2,
        "cost_price": 9000,
        "sale_price": 15000,
        "supplier": "Toshkent Sport",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["ogil bolalar trikosi", "adidas bolalar trikosi", "kichik triko", "5 yosh triko"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 23,
        "sku": "MS-1023",
        "name": "1-2 yoshli o'g'il bolalar trikosi",
        "category": "Shim",
        "brand": "MarkazSavdo",
        "sizes": ["1 yosh", "2 yosh"],
        "colors": ["moviy", "kamuflyaj"],
        "gender": "Bolalar",
        "material": "Paxta trikotaj",
        "stock": 25,
        "min_stock": 3,
        "cost_price": 5000,
        "sale_price": 8000,
        "supplier": "Namangan Baby",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["chaqaloq trikosi", "1 yosh triko", "2 yosh triko", "kichkintoy trikosi"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 24,
        "sku": "MS-1024",
        "name": "2 kishilik pastel jild komplekti",
        "category": "Uy tekstili",
        "brand": "MarkazSavdo",
        "sizes": ["2 kishilik"],
        "colors": ["pushti", "kulrang"],
        "gender": "Uniseks",
        "material": "Paxta / Saten",
        "stock": 15,
        "min_stock": 2,
        "cost_price": 38000,
        "sale_price": 60000,
        "supplier": "Buxoro Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["pastel jild", "korpa jildi", "yostiq jildi", "choyshab komplekti", "to'shak jildi"],
        "image_url": "https://images.unsplash.com/photo-1522771739844-6a9f6d5f14af?w=600"
    },
    {
        "id": 25,
        "sku": "MS-1025",
        "name": "Ayollar qora kardigan kostyum",
        "category": "Kardigan",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["qora"],
        "gender": "Ayol",
        "material": "Jun aralash trikotaj",
        "stock": 1,
        "min_stock": 2,
        "cost_price": 105000,
        "sale_price": 160000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Kam qolgan",
        "aliases": ["kardigan kostyum", "qora kardigan", "ayollar kardigani", "kardigan"],
        "image_url": "https://images.unsplash.com/photo-1585487000160-6ebcfceb0d03?w=600"
    },
    {
        "id": 26,
        "sku": "MS-1026",
        "name": "Qizlar uchun katakli kardigan",
        "category": "Kardigan",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["qora-oq", "katakli"],
        "gender": "Ayol",
        "material": "Yumshoq trikotaj",
        "stock": 7,
        "min_stock": 2,
        "cost_price": 75000,
        "sale_price": 115000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["katakli kardigan", "qizlar kardigani", "katak kardigan"],
        "image_url": "https://images.unsplash.com/photo-1585487000160-6ebcfceb0d03?w=600"
    },
    {
        "id": 27,
        "sku": "MS-1027",
        "name": "Ayollar sutrang oq tonika",
        "category": "Ko'ylak",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["sutrang oq", "oq"],
        "gender": "Ayol",
        "material": "Viskoza / Paxta",
        "stock": 10,
        "min_stock": 2,
        "cost_price": 52000,
        "sale_price": 80000,
        "supplier": "Samarqand Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["sutrang tonika", "oq tonika", "ayollar tonikasi", "tonika"],
        "image_url": "https://images.unsplash.com/photo-1595777457583-95e059d581b8?w=600"
    },
    {
        "id": 28,
        "sku": "MS-1028",
        "name": "Ayollar to'q rangli nafis tonika",
        "category": "Ko'ylak",
        "brand": "MarkazSavdo",
        "sizes": ["Standart"],
        "colors": ["to'q kulrang", "qora"],
        "gender": "Ayol",
        "material": "Sifatli trikotaj",
        "stock": 8,
        "min_stock": 2,
        "cost_price": 95000,
        "sale_price": 150000,
        "supplier": "Samarqand Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["nafis tonika", "toq tonika", "qishki tonika", "ayollar to'q tonikasi"],
        "image_url": "https://images.unsplash.com/photo-1595777457583-95e059d581b8?w=600"
    },
    {
        "id": 29,
        "sku": "MS-1029",
        "name": "Erkaklar ikki tomonlama vitrofka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["52", "54", "56"],
        "colors": ["yashil", "qora"],
        "gender": "Erkak",
        "material": "Ikki tomonlama plashovka",
        "stock": 11,
        "min_stock": 2,
        "cost_price": 80000,
        "sale_price": 120000,
        "supplier": "Xitoy Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["ikki tomonlama vitrofka", "yashil vitrofka", "dvuxstoronniy vitrofka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 30,
        "sku": "MS-1030",
        "name": "Erkaklar to'q ko'k vitrofka",
        "category": "Kurtka",
        "brand": "MarkazSavdo",
        "sizes": ["52", "54", "56"],
        "colors": ["to'q ko'k"],
        "gender": "Erkak",
        "material": "Plashovka",
        "stock": 13,
        "min_stock": 2,
        "cost_price": 90000,
        "sale_price": 140000,
        "supplier": "Xitoy Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["to'q ko'k vitrofka", "kurtka toq kok", "shamol otkazmas"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 31,
        "sku": "MS-1031",
        "name": "Yosh bolalar Boss vitrofka",
        "category": "Kurtka",
        "brand": "Boss",
        "sizes": ["Standart", "5-8 yosh"],
        "colors": ["to'q ko'k"],
        "gender": "Bolalar",
        "material": "Plashovka",
        "stock": 9,
        "min_stock": 2,
        "cost_price": 55000,
        "sale_price": 85000,
        "supplier": "Guangzhou Baby",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["boss vitrofka", "bolalar boss", "bolalar vitrofkasi", "boss kurtka"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 32,
        "sku": "MS-1032",
        "name": "Bolalar jinsi troyka komplekt",
        "category": "Kostyum",
        "brand": "MarkazSavdo",
        "sizes": ["90", "100", "110"],
        "colors": ["bej", "ko'k"],
        "gender": "Bolalar",
        "material": "Jinsi + paxta",
        "stock": 6,
        "min_stock": 2,
        "cost_price": 110000,
        "sale_price": 170000,
        "supplier": "Turkiya Baby",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["jinsi troyka", "troyka komplekt", "bolalar troykasi", "bolalar jinsi komplekti"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 33,
        "sku": "MS-1033",
        "name": "Yosh bolalar Polo vitrofka",
        "category": "Kurtka",
        "brand": "Polo",
        "sizes": ["130", "140", "150", "160", "170"],
        "colors": ["to'q ko'k"],
        "gender": "Bolalar",
        "material": "Plashovka",
        "stock": 10,
        "min_stock": 2,
        "cost_price": 62000,
        "sale_price": 95000,
        "supplier": "Guangzhou Baby",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["polo vitrofka bolalar", "bolalar polo kurtka", "polo bolalar"],
        "image_url": "https://images.unsplash.com/photo-1544441893-675973e31985?w=600"
    },
    {
        "id": 34,
        "sku": "MS-1034",
        "name": "Erkaklar AQSh bayroqli svitir",
        "category": "Svitir",
        "brand": "MarkazSavdo",
        "sizes": ["L", "XL", "XXL"],
        "colors": ["xaki", "zaytun"],
        "gender": "Erkak",
        "material": "Paxta trikotaj",
        "stock": 8,
        "min_stock": 2,
        "cost_price": 45000,
        "sale_price": 70000,
        "supplier": "Toshkent Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["aqsh bayroqli svitir", "xaki svitir", "bayroqli kofta", "amerika svitir"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 35,
        "sku": "MS-1035",
        "name": "Erkaklar kapyushonli xudi svitir",
        "category": "Svitir",
        "brand": "MarkazSavdo",
        "sizes": ["L", "XL", "XXL"],
        "colors": ["kulrang"],
        "gender": "Erkak",
        "material": "Qalin trikotaj",
        "stock": 12,
        "min_stock": 2,
        "cost_price": 55000,
        "sale_price": 85000,
        "supplier": "Toshkent Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["kapyushonli xudi", "xudi svitir", "kulrang xudi", "erkaklar xudisi"],
        "image_url": "https://images.unsplash.com/photo-1620799140408-edc6dcb6d633?w=600"
    },
    {
        "id": 36,
        "sku": "MS-1036",
        "name": "Erkaklar kulrang sportivka kostyum",
        "category": "Kostyum",
        "brand": "MarkazSavdo",
        "sizes": ["L", "XL", "XXL"],
        "colors": ["kulrang"],
        "gender": "Erkak",
        "material": "Sifatli ikki ipli trikotaj",
        "stock": 7,
        "min_stock": 2,
        "cost_price": 140000,
        "sale_price": 210000,
        "supplier": "Turkiya Import",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["kulrang sportivka", "sportivka to'plam", "masterka shim", "sport kostyumi"],
        "image_url": "https://images.unsplash.com/photo-1515886657613-9f3515b0c78f?w=600"
    },
    {
        "id": 37,
        "sku": "MS-1037",
        "name": "Ayollar dvoyka ko'ylak-yubka komplekt",
        "category": "Ko'ylak",
        "brand": "MarkazSavdo",
        "sizes": ["48", "50", "52", "54"],
        "colors": ["bordo", "qora", "kulrang"],
        "gender": "Ayol",
        "material": "Trikotaj va jakkard",
        "stock": 8,
        "min_stock": 2,
        "cost_price": 145000,
        "sale_price": 220000,
        "supplier": "Samarqand Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["dvoyka koylak yubka", "koylak yubka komplekt", "ayollar dvoykasi", "ayollar komplekt"],
        "image_url": "https://images.unsplash.com/photo-1585487000160-6ebcfceb0d03?w=600"
    },
    {
        "id": 38,
        "sku": "MS-1038",
        "name": "Qiz bolalar uchun pijama komplekt",
        "category": "Pijama",
        "brand": "MarkazSavdo",
        "sizes": ["Standart", "4-7 yosh"],
        "colors": ["siyohrang", "sariq", "mentol"],
        "gender": "Bolalar",
        "material": "Paxta 100%",
        "stock": 15,
        "min_stock": 3,
        "cost_price": 15000,
        "sale_price": 25000,
        "supplier": "Namangan Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["qizlar pijamasi", "pijama komplekt", "uy pijamasi", "bolalar pijamasi"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 39,
        "sku": "MS-1039",
        "name": "O'g'il bolalar uchun pijama komplekt",
        "category": "Pijama",
        "brand": "MarkazSavdo",
        "sizes": ["Standart", "4-7 yosh"],
        "colors": ["havorang", "qora", "yashil"],
        "gender": "Bolalar",
        "material": "Paxta 100%",
        "stock": 14,
        "min_stock": 3,
        "cost_price": 18000,
        "sale_price": 30000,
        "supplier": "Namangan Tekstil",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["ogil bolalar pijamasi", "ogil bola pijama", "bolalar uy kiyimi"],
        "image_url": "https://images.unsplash.com/photo-1519725392576-96a84f3eb48c?w=600"
    },
    {
        "id": 40,
        "sku": "MS-1040",
        "name": "O'g'il bolalar Xitoy krossovka",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["26", "27", "28", "29", "30"],
        "colors": ["qora-oq"],
        "gender": "Bolalar",
        "material": "To'qima mato va rezina",
        "stock": 12,
        "min_stock": 2,
        "cost_price": 48000,
        "sale_price": 75000,
        "supplier": "Guangzhou Shoes",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["ogil bola krossovka", "xitoy krossovka", "bolalar krossovkasi", "qora oq krossovka"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    },
    {
        "id": 41,
        "sku": "MS-1041",
        "name": "Qiz bolalar Xitoy krossovka",
        "category": "Oyoq kiyim",
        "brand": "MarkazSavdo",
        "sizes": ["26", "27", "28", "29", "30", "31"],
        "colors": ["oq-binafsharang", "pushti"],
        "gender": "Bolalar",
        "material": "To'qima mato va yengil taglik",
        "stock": 10,
        "min_stock": 2,
        "cost_price": 55000,
        "sale_price": 85000,
        "supplier": "Guangzhou Shoes",
        "incoming_date": "04.10.2026",
        "status": "Mavjud",
        "aliases": ["qizlar krossovkasi", "oq binafsha krossovka", "pushti krossovka", "bolalar krossovkalari"],
        "image_url": "https://images.unsplash.com/photo-1549298916-b41d501d3772?w=600"
    }
]

def build_inventory():
    print(f"Jami {len(REAL_PRODUCTS)} ta haqiqiy do'kon tovari qayta ishlanmoqda...")

    # 1. products.json faylini yaratish
    json_path = ROOT_DIR / "products.json"
    json_data = []
    for p in REAL_PRODUCTS:
        json_data.append({
            "id": p["id"],
            "sku": p["sku"],
            "name": p["name"],
            "category": p["category"],
            "brand": p["brand"],
            "price": p["sale_price"],
            "cost_price": p["cost_price"],
            "sizes": p["sizes"],
            "colors": [c.lower() for c in p["colors"]],
            "gender": p["gender"],
            "material": p["material"],
            "stock": p["stock"],
            "min_stock": p["min_stock"],
            "supplier": p["supplier"],
            "status": p["status"],
            "aliases": p["aliases"],
            "image_url": p["image_url"]
        })

    # Write products.json atomically
    atomic_write_json(str(json_path), json_data, indent=2)
    print(f"✅ products.json muvaffaqiyatli saqlandi ({len(json_data)} haqiqiy tovar)!")

    # 2. SQLite bazasini to'ldirish
    db_path = ROOT_DIR / "ingichka_store.db"
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    cursor.execute("PRAGMA table_info(products)")
    existing_cols = [c[1] for c in cursor.fetchall()]
    new_cols = [
        ("sku", "TEXT"),
        ("brand", "TEXT"),
        ("gender", "TEXT"),
        ("material", "TEXT"),
        ("min_stock", "INTEGER DEFAULT 0"),
        ("supplier", "TEXT")
    ]
    for col_name, col_type in new_cols:
        if col_name not in existing_cols:
            cursor.execute(f"ALTER TABLE products ADD COLUMN {col_name} {col_type};")

    # Baza tovarlarini tozalash va haqiqiy tovarlar bilan to'ldirish
    cursor.execute("DELETE FROM products;")
    for p in REAL_PRODUCTS:
        desc = f"Brend: {p['brand']}, Jinsi: {p['gender']}, Material: {p['material']}, SKU: {p['sku']}. " + ", ".join(p["aliases"])
        sizes_str = ", ".join(p["sizes"])
        colors_str = ", ".join(p["colors"])
        cursor.execute("""
            INSERT INTO products (
                id, name, category, size, color, cost_price, sale_price, stock_quantity,
                description, is_active, sku, brand, gender, material, min_stock, supplier
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?, ?)
        """, (
            p["id"], p["name"], p["category"], sizes_str, colors_str,
            p["cost_price"], p["sale_price"], p["stock"],
            desc, p["sku"], p["brand"], p["gender"], p["material"], p["min_stock"], p["supplier"]
        ))
    conn.commit()
    conn.close()
    print(f"✅ SQLite bazasi (ingichka_store.db) {len(REAL_PRODUCTS)} ta real tovar bilan 100% to'ldirildi!")

    # 3. Professional Excel hisobotini yaratish
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ombor Hisoboti"

        # Sarlavha
        ws.merge_cells("A1:P1")
        title_cell = ws["A1"]
        title_cell.value = "MARKAZSAVDO DO'KONI — HAQIQIY OMBOR HISOBOTI"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:P2")
        sub_cell = ws["A2"]
        sub_cell.value = "Yaratilgan sana: 05.10.2026 | Kanal: @markazsavdo_7 | Do'kon egasi: Mansur Radjabov"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = [
            "№", "SKU", "Mahsulot nomi", "Kategoriya", "Brend", "O'lcham", "Rang",
            "Jins", "Material", "Miqdor", "Min. qoldiq", "Sotib olish narxi (so'm)",
            "Sotish narxi (so'm)", "Yetkazib beruvchi", "Kirim sanasi", "Holat"
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        status_fill_kam = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Sariq
        status_fill_tug = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Qizilroq
        status_fill_mav = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Yashil

        for row_idx, p in enumerate(REAL_PRODUCTS, 5):
            ws.cell(row=row_idx, column=1, value=p["id"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=p["sku"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=p["name"]).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=4, value=p["category"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=5, value=p["brand"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=6, value=", ".join(p["sizes"])).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=7, value=", ".join(p["colors"])).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=8, value=p["gender"]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=9, value=p["material"]).alignment = Alignment(horizontal="center")

            # Miqdor va narxlar
            c_stock = ws.cell(row=row_idx, column=10, value=p["stock"])
            c_stock.alignment = Alignment(horizontal="right")
            c_stock.number_format = "#,##0"

            c_min = ws.cell(row=row_idx, column=11, value=p["min_stock"])
            c_min.alignment = Alignment(horizontal="right")
            c_min.number_format = "#,##0"

            c_cost = ws.cell(row=row_idx, column=12, value=p["cost_price"])
            c_cost.alignment = Alignment(horizontal="right")
            c_cost.number_format = "#,##0"

            c_sale = ws.cell(row=row_idx, column=13, value=p["sale_price"])
            c_sale.alignment = Alignment(horizontal="right")
            c_sale.number_format = "#,##0"

            ws.cell(row=row_idx, column=14, value=p["supplier"]).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=15, value=p["incoming_date"]).alignment = Alignment(horizontal="center")

            c_stat = ws.cell(row=row_idx, column=16, value=p["status"])
            c_stat.alignment = Alignment(horizontal="center")

            if p["status"] == "Kam qolgan":
                c_stat.fill = status_fill_kam
            elif p["status"] == "Tugagan":
                c_stat.fill = status_fill_tug
            else:
                c_stat.fill = status_fill_mav

            for col_i in range(1, 17):
                ws.cell(row=row_idx, column=col_i).border = border_thin

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        excel_path = ROOT_DIR / "reports" / "KIYIM_KECHAK_DOKONI_OMBOR_HISOBOTI.xlsx"
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(excel_path))
        print(f"✅ Yangilangan professional Excel fayli saqlandi: {excel_path}!")
    except Exception as e:
        print(f"Excel faylini yaratishda xato: {e}")

if __name__ == "__main__":
    build_inventory()
