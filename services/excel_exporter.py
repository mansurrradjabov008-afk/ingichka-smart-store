import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path
from config import DB_PATH, STORE_NAME
from utils.file_utils import atomic_write_binary

REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

class ExcelExporter:
    """
    Do'kon egasi uchun professional moliyaviy va ombor hisobotlarini
    Excel (.xlsx) formatida tayyorlab berish.
    """

    @staticmethod
    def generate_full_report() -> str:
        wb = openpyxl.Workbook()
        
        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        bold_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        # ----------------------------------------------------
        # SHEET 1: BUYURTMALAR VA SAVDO
        # ----------------------------------------------------
        ws_orders = wb.active
        ws_orders.title = "Buyurtmalar Tarixi"

        ws_orders.merge_cells("A1:G1")
        ws_orders["A1"] = f"{STORE_NAME} - Barcha Buyurtmalar Hisoboti"
        ws_orders["A1"].font = title_font
        ws_orders["A1"].alignment = left_align

        ws_orders["A2"] = f"Hisobot yaratilgan vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws_orders["A2"].font = Font(name="Calibri", size=9, italic=True)

        order_headers = ["ID", "Sana va Vaqt", "Xaridor", "Telefon", "Manzil", "Summa (so'm)", "To'lov turi", "Holati"]
        for col_idx, h in enumerate(order_headers, 1):
            cell = ws_orders.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        cursor.execute("SELECT * FROM orders ORDER BY id DESC")
        orders = cursor.fetchall()
        total_revenue = 0.0

        for r_idx, o in enumerate(orders, 5):
            ws_orders.cell(row=r_idx, column=1, value=f"#{o['id']}").alignment = center_align
            ws_orders.cell(row=r_idx, column=2, value=str(o['created_at'])).alignment = center_align
            ws_orders.cell(row=r_idx, column=3, value=o['customer_name']).alignment = left_align
            ws_orders.cell(row=r_idx, column=4, value=o['customer_phone']).alignment = center_align
            ws_orders.cell(row=r_idx, column=5, value=o['delivery_address']).alignment = left_align
            
            sum_cell = ws_orders.cell(row=r_idx, column=6, value=o['total_amount'])
            sum_cell.number_format = '#,##0'
            sum_cell.alignment = right_align
            total_revenue += o['total_amount']

            ws_orders.cell(row=r_idx, column=7, value=o['payment_method']).alignment = center_align
            ws_orders.cell(row=r_idx, column=8, value=o['status'].capitalize()).alignment = center_align

            for c in range(1, 9):
                ws_orders.cell(row=r_idx, column=c).border = thin_border

        # Jami tushum qatori
        tot_row = len(orders) + 5
        ws_orders.cell(row=tot_row, column=5, value="JAMI TUSHUM:").font = bold_font
        ws_orders.cell(row=tot_row, column=5).alignment = right_align
        tot_sum_cell = ws_orders.cell(row=tot_row, column=6, value=total_revenue)
        tot_sum_cell.font = bold_font
        tot_sum_cell.number_format = '#,##0'
        tot_sum_cell.alignment = right_align

        # ----------------------------------------------------
        # SHEET 2: OMBOR VA TOVAR QOLDIG'I
        # ----------------------------------------------------
        ws_inventory = wb.create_sheet(title="Ombor Qoldig'i")

        ws_inventory.merge_cells("A1:H1")
        ws_inventory["A1"] = f"{STORE_NAME} - Ombor Holati va Qoldiqlari"
        ws_inventory["A1"].font = title_font

        inv_headers = ["ID", "Tovar Nomi", "Kategoriya", "O'lcham", "Rang", "Tan Narx", "Sotuv Narx", "Omborda (dona)", "Ombor Qiymati"]
        for col_idx, h in enumerate(inv_headers, 1):
            cell = ws_inventory.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        cursor.execute("SELECT * FROM products WHERE is_active = 1 ORDER BY category, name")
        products = cursor.fetchall()
        total_stock_value = 0.0

        for r_idx, p in enumerate(products, 4):
            ws_inventory.cell(row=r_idx, column=1, value=f"#{p['id']}").alignment = center_align
            ws_inventory.cell(row=r_idx, column=2, value=p['name']).alignment = left_align
            ws_inventory.cell(row=r_idx, column=3, value=p['category']).alignment = left_align
            ws_inventory.cell(row=r_idx, column=4, value=p['size']).alignment = center_align
            ws_inventory.cell(row=r_idx, column=5, value=p['color']).alignment = center_align
            
            c_cost = ws_inventory.cell(row=r_idx, column=6, value=p['cost_price'])
            c_cost.number_format = '#,##0'
            c_cost.alignment = right_align

            c_sale = ws_inventory.cell(row=r_idx, column=7, value=p['sale_price'])
            c_sale.number_format = '#,##0'
            c_sale.alignment = right_align

            c_qty = ws_inventory.cell(row=r_idx, column=8, value=p['stock_quantity'])
            c_qty.alignment = center_align
            if p['stock_quantity'] <= 5:
                c_qty.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Red alert for low stock

            val = p['stock_quantity'] * p['sale_price']
            c_val = ws_inventory.cell(row=r_idx, column=9, value=val)
            c_val.number_format = '#,##0'
            c_val.alignment = right_align
            total_stock_value += val

            for c in range(1, 10):
                ws_inventory.cell(row=r_idx, column=c).border = thin_border

        inv_tot_row = len(products) + 4
        ws_inventory.cell(row=inv_tot_row, column=8, value="JAMI QIYMAT:").font = bold_font
        ws_inventory.cell(row=inv_tot_row, column=8).alignment = right_align
        val_sum_cell = ws_inventory.cell(row=inv_tot_row, column=9, value=total_stock_value)
        val_sum_cell.font = bold_font
        val_sum_cell.number_format = '#,##0'
        val_sum_cell.alignment = right_align

        # Auto-fit column widths
        for ws in [ws_orders, ws_inventory]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        conn.close()

        # Save file atomically (Rule 4)
        file_path = REPORTS_DIR / f"Ingichka_Kassa_Hisoboti_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        atomic_write_binary(str(file_path), lambda tmp: wb.save(tmp))
        return str(file_path)
